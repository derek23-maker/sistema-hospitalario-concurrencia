"""Genera la guia de defensa del Sistema Hospitalario en formato Word."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parent / "Guia_para_defender_Sistema_Hospitalario_Sprint_2.docx"

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRAY) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def set_run_font(run, size: float, bold: bool = False, color: RGBColor = BLACK, name: str = "Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(document: Document, text: str, *, bold_lead: str | None = None, size: float = 10.2) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size, bold=True)
        remainder = paragraph.add_run(text[len(bold_lead):])
        set_run_font(remainder, size)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, 9.7)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(4 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    run = paragraph.add_run(text)
    set_run_font(run, 15 if level == 1 else 11.2, bold=True)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, 9.2, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            set_run_font(run, 8.8)
            if widths:
                cell.width = Inches(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_question(document: Document, question: str, answer: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.06
    question_run = paragraph.add_run(question + " ")
    set_run_font(question_run, 9.2, bold=True)
    answer_run = paragraph.add_run(answer)
    set_run_font(answer_run, 9.2)


def page_break(document: Document) -> None:
    document.add_page_break()


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = BLACK

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = BLACK

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("Sistema Hospitalario  Sprint 2")
    set_run_font(header_run, 8.3, color=RGBColor(70, 70, 70))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer_run = footer.add_run("Guia de defensa  Pagina ")
    set_run_font(footer_run, 8.2, color=RGBColor(70, 70, 70))
    add_page_field(footer)

    document.core_properties.title = "Guia para defender el Sistema Hospitalario de Gestion de Citas"
    document.core_properties.subject = "Guia de estudio y defensa oral del Sprint 2"
    document.core_properties.author = "Equipo Sistema Hospitalario"


def build_document() -> Document:
    document = Document()
    configure_document(document)

    # Pagina 1
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(55)
    title.paragraph_format.space_after = Pt(12)
    title_run = title.add_run("Guia para defender el Sistema Hospitalario de Gestion de Citas")
    set_run_font(title_run, 23, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    subtitle_run = subtitle.add_run("Sprint 2  Programacion Concurrente")
    set_run_font(subtitle_run, 14, bold=True, color=RGBColor(70, 70, 70))

    add_text(document, "Esta guia explica, con palabras directas, que hace cada parte del proyecto y como se relacionan entre si. Sirve para estudiar antes de la presentacion y para responder preguntas del profesor sin memorizar codigo linea por linea.")
    add_text(document, "El sistema permite consultar horarios medicos, reservar una cita y cancelarla. La idea importante del Sprint 2 no es solo mostrar una pantalla: es demostrar que dos personas no pueden reservar al mismo medico en el mismo horario cuando usan el sistema al mismo tiempo.")
    add_heading(document, "El recorrido de una cita", 2)
    add_table(
        document,
        ["Paso", "Componente", "Que ocurre"],
        [
            ["1", "Flet", "El usuario escribe los identificadores, fecha y motivo. La pantalla pide los horarios disponibles."],
            ["2", "FastAPI", "La API revisa los datos y aplica las reglas del sistema. Flet no toca MySQL directamente."],
            ["3", "MySQL", "La base consulta horarios, guarda la cita y protege la reserva con una transaccion."],
            ["4", "Respuesta", "FastAPI comunica exito, error de validacion o conflicto. Flet lo muestra al usuario."],
        ],
        [0.45, 1.15, 4.95],
    )
    add_text(document, "La frase corta para presentar la arquitectura es esta: Flet es la interfaz, FastAPI es el intermediario con reglas y MySQL guarda los datos. El procedimiento almacenado de MySQL decide finalmente si una cita se puede confirmar.")
    page_break(document)

    # Pagina 2
    add_heading(document, "1 Arquitectura del sistema")
    add_text(document, "Una arquitectura es la forma de dividir un programa para que cada parte tenga una responsabilidad clara. En este proyecto se separaron la pantalla, la logica del negocio y los datos. Esta separacion facilita corregir una parte sin reescribir las demas.")
    add_table(
        document,
        ["Capa", "Responsabilidad", "Ejemplo en el proyecto"],
        [
            ["Frontend", "Recibir acciones y mostrar resultados", "La aplicacion Flet donde se elige un horario y se presiona Reservar cita"],
            ["Backend", "Validar solicitudes y coordinar reglas", "FastAPI con las rutas horarios reservar y cancelar"],
            ["Base de datos", "Guardar informacion de forma permanente", "MySQL con doctores pacientes servicios horarios y citas"],
            ["Procedimiento", "Proteger una operacion critica", "sp_reservar_cita usa transaccion y bloqueo pesimista"],
        ],
        [1.05, 2.45, 3.05],
    )
    add_heading(document, "Por que no conectar Flet directamente con MySQL", 2)
    add_text(document, "Si la pantalla tuviera usuario y contraseña de MySQL, cualquier persona podria inspeccionar o modificar tablas desde la aplicacion. Tambien se repetiria la logica de validacion en cada pantalla. FastAPI evita ese problema porque concentra la comunicacion con la base y decide que operaciones estan permitidas.")
    add_text(document, "La API funciona como una recepcion. Recibe una solicitud, revisa si tiene los datos necesarios, consulta la base y devuelve una respuesta. La pantalla solo necesita saber si hay horarios, si la cita se confirmo o si alguien tomo el turno antes.")
    add_heading(document, "Pregunta probable", 2)
    add_question(document, "Por que se usan tres capas", "Porque cada capa resuelve una tarea distinta. La interfaz no debe administrar la base de datos y la base no debe decidir como se dibuja una pantalla. Separarlas reduce errores y permite cambiar una capa con menor impacto en las otras.")
    add_question(document, "Que pasa si FastAPI se apaga", "Flet no puede completar nuevas operaciones porque su intermediario no responde. Los datos que ya estaban guardados siguen en MySQL; el problema es de disponibilidad de la API, no de perdida inmediata de la informacion.")
    page_break(document)

    # Pagina 3
    add_heading(document, "2 Frontend con Flet")
    add_text(document, "El frontend es la parte visible para la persona usuaria. En este sistema es una aplicacion hecha con Flet, una libreria de Python que permite crear interfaces de escritorio o web con controles como campos de texto, botones y listas desplegables.")
    add_text(document, "La pantalla solicita tres identificadores: paciente, medico y servicio. Tambien recibe una fecha y un motivo de consulta. Cuando el usuario presiona Consultar horarios, Flet manda una solicitud a la API. La API responde con horarios libres y ocupados, y Flet llena la lista para que el usuario elija una hora disponible.")
    add_heading(document, "Que significa asincrono", 2)
    add_text(document, "Una operacion asincrona permite que la interfaz espere una respuesta de red sin congelarse. Mientras FastAPI y MySQL trabajan, Flet puede mostrar un indicador de carga. En el codigo se usan funciones async para consultar horarios, reservar y cancelar.")
    add_table(
        document,
        ["Accion en Flet", "Solicitud enviada", "Resultado visible"],
        [
            ["Consultar horarios", "GET a horarios", "Lista de horas disponibles o mensaje de ausencia"],
            ["Reservar cita", "POST a citas reservar", "Confirmacion o aviso de conflicto"],
            ["Cancelar cita", "PATCH a citas cancelar", "Estado de cita cancelada"],
        ],
        [1.55, 2.2, 2.8],
    )
    add_text(document, "Flet no guarda por si mismo una cita. Solo muestra controles y consume la API. Si se cerrara Flet despues de una reserva exitosa, la cita continuaria guardada porque la informacion vive en MySQL.")
    add_heading(document, "Preguntas probables", 2)
    add_question(document, "Flet es una base de datos", "No. Flet sirve para construir la interfaz. La base de datos es MySQL.")
    add_question(document, "Por que no aparece una pantalla para crear doctores", "El alcance del Sprint 2 se concentra en agenda, reservacion y cancelacion. Los doctores, pacientes, servicios y horarios se cargan inicialmente con MySQL. Un modulo administrativo puede agregarse despues.")
    page_break(document)

    # Pagina 4
    add_heading(document, "3 Backend con FastAPI")
    add_text(document, "El backend es el programa que atiende solicitudes de la interfaz y aplica las reglas antes de tocar la base de datos. En este proyecto se usa FastAPI, un framework de Python para crear una API REST. REST es una forma comun de organizar servicios web mediante rutas, metodos HTTP y respuestas en formato JSON.")
    add_text(document, "FastAPI recibe una peticion como reservar una cita, convierte los datos a tipos de Python, valida que los identificadores sean positivos y que la fecha y hora tengan un formato correcto. Despues envia la solicitud al servicio de agenda, que consulta el repositorio de MySQL.")
    add_heading(document, "Partes internas del backend", 2)
    add_table(
        document,
        ["Carpeta", "Funcion", "Ejemplo"],
        [
            ["api routes", "Define las direcciones que recibe la API", "POST citas reservar"],
            ["services", "Contiene reglas de negocio", "Revisar que la cita cabe dentro de la jornada"],
            ["repositories", "Hace consultas a MySQL", "Buscar horario del medico o llamar al procedimiento"],
            ["domain", "Define datos esperados y resultados", "ReservaCitaRequest y HorariosResponse"],
            ["infrastructure", "Configura conexion tecnica", "Pool de conexiones MySQL"],
        ],
        [1.25, 2.25, 3.05],
    )
    add_text(document, "Esta separacion es util en una defensa porque muestra que el codigo no esta todo mezclado en un solo archivo. La ruta recibe la peticion, el servicio decide las reglas y el repositorio se comunica con la base.")
    add_heading(document, "Validacion y respuestas", 2)
    add_text(document, "FastAPI responde con codigos HTTP. Un 201 significa que se creo una cita. Un 404 indica que no existe un paciente, medico o servicio. Un 409 significa conflicto, por ejemplo cuando otra solicitud gano el mismo horario. Un 422 indica datos validos en forma pero imposibles para la regla, como reservar fuera de la jornada.")
    page_break(document)

    # Pagina 5
    add_heading(document, "4 Base de datos MySQL")
    add_text(document, "MySQL es el sistema que almacena la informacion de forma permanente. Las tablas organizan datos relacionados. Cada fila representa un registro y cada columna representa un atributo. Por ejemplo, la tabla pacientes guarda el nombre, telefono y fecha de nacimiento de cada paciente.")
    add_text(document, "La base se usa con el motor InnoDB. InnoDB importa en esta materia porque soporta transacciones, bloqueos por fila y llaves foraneas. Estas caracteristicas hacen posible reservar una cita sin que dos usuarios confirmen el mismo espacio al mismo tiempo.")
    add_heading(document, "Tablas que participan en la demostracion", 2)
    add_table(
        document,
        ["Tabla", "Que guarda", "Relacion principal"],
        [
            ["cuentas_usuario", "Datos de acceso y rol", "Una cuenta puede corresponder a paciente o doctor"],
            ["pacientes", "Datos personales del paciente", "Cada paciente se usa al crear una cita"],
            ["doctores", "Datos del medico y especialidad", "Un doctor tiene uno o varios horarios"],
            ["servicios", "Nombre costo y duracion", "La duracion define el tamano del turno"],
            ["horarios_medicos", "Jornada semanal del doctor", "Indica cuando puede atender"],
            ["citas", "Reservas y estados", "Une paciente doctor servicio fecha y hora"],
            ["ausencias_medicos", "Vacaciones permisos o incapacidad", "Bloquea horarios durante una ausencia aprobada"],
        ],
        [1.35, 2.45, 2.75],
    )
    add_heading(document, "Llaves y relaciones", 2)
    add_text(document, "Una llave primaria identifica una fila de manera unica, por ejemplo id_doctor. Una llave foranea conecta tablas, por ejemplo citas.id_doctor apunta a doctores.id_doctor. Gracias a estas relaciones MySQL evita guardar una cita asociada a un doctor inexistente.")
    add_text(document, "Los datos iniciales no aparecen solos. El archivo 02_CARGAR_DATOS_DEMO.sql agrega un paciente, un doctor, una especialidad, un servicio y un horario. Las citas posteriores se crean desde Flet a traves de FastAPI.")
    page_break(document)

    # Pagina 6
    add_heading(document, "5 Concurrencia y condiciones de carrera")
    add_text(document, "Concurrencia significa que varias tareas pueden avanzar al mismo tiempo. En este sistema ocurre cuando dos personas usan la aplicacion y tratan de reservar. El recurso compartido es la agenda de un medico en una fecha y hora especificas.")
    add_heading(document, "El problema que se debe evitar", 2)
    add_text(document, "Supongamos que Ana y Luis consultan el horario de las 10:00. Ambos ven que esta libre. Si los dos presionan Reservar al mismo tiempo y el sistema solo hiciera una consulta simple, ambos podrian intentar insertar la misma cita. Eso se llama condicion de carrera: el resultado depende de cual peticion llega primero.")
    add_table(
        document,
        ["Sin proteccion", "Con transaccion y bloqueo"],
        [
            ["Ana y Luis leen que 10:00 esta libre. Los dos intentan guardar. Puede haber doble reserva o un error dificil de explicar.", "Ana toma un bloqueo de la agenda del medico. Luis espera. Cuando Ana confirma, Luis vuelve a revisar y recibe conflicto porque el horario ya esta ocupado."],
        ],
        [3.3, 3.3],
    )
    add_heading(document, "Conceptos que deben poder explicar", 2)
    add_text(document, "Una transaccion agrupa pasos que deben cumplirse juntos. Si todo sale bien se ejecuta COMMIT y los cambios quedan guardados. Si hay una regla incumplida o un error se ejecuta ROLLBACK y los cambios parciales se deshacen.")
    add_text(document, "SELECT FOR UPDATE es un bloqueo pesimista. Pesimista no significa que el sistema falle; significa que asume que puede haber competencia y bloquea el recurso antes de modificarlo. En nuestro caso se bloquea la fila del doctor durante la reserva para ordenar las solicitudes de ese medico.")
    add_question(document, "Por que no basta con que Flet o FastAPI revisen si el horario esta libre", "Porque esa revision puede quedar vieja en milisegundos. La base de datos es el unico lugar compartido por todas las peticiones, por eso la decision final se protege dentro de MySQL.")
    page_break(document)

    # Pagina 7
    add_heading(document, "6 Procedimiento almacenado de reserva")
    add_text(document, "Un procedimiento almacenado es codigo SQL guardado dentro de MySQL. FastAPI llama a sp_reservar_cita con el paciente, medico, servicio, fecha, hora y motivo. La base responde con un texto de resultado que FastAPI traduce a una respuesta para Flet.")
    add_heading(document, "Que hace sp_reservar_cita", 2)
    add_table(
        document,
        ["Paso", "Regla aplicada"],
        [
            ["1", "Inicia una transaccion."],
            ["2", "Comprueba que existen el paciente, el medico y el servicio."],
            ["3", "Bloquea al medico con SELECT FOR UPDATE para serializar reservas de su agenda."],
            ["4", "Comprueba que la hora cabe dentro de la jornada semanal del medico."],
            ["5", "Comprueba que no existe una ausencia aprobada que cubra ese turno."],
            ["6", "Busca traslapes con citas pendientes o confirmadas."],
            ["7", "Inserta una nueva cita o reutiliza la fila de una cita cancelada."],
            ["8", "Hace COMMIT si se confirma o ROLLBACK si existe un problema."],
        ],
        [0.55, 6.05],
    )
    add_text(document, "La llave unica uk_doctor_fecha_hora es una segunda defensa. No permite dos filas con el mismo doctor, fecha y hora. El procedimiento corregido tambien reutiliza un registro cancelado para que ese horario se pueda reservar otra vez sin romper la llave unica.")
    add_heading(document, "Resultados esperados", 2)
    add_table(
        document,
        ["Resultado de MySQL", "Respuesta de la API", "Significado"],
        [
            ["OK APPOINTMENT CONFIRMED", "201", "La cita quedo confirmada"],
            ["CONFLICT SLOT TAKEN", "409", "Otro usuario tomo la hora o se detecto traslape"],
            ["CONFLICT DOCTOR ABSENT", "409", "El medico esta ausente"],
            ["VALIDATION OUTSIDE WORKING HOURS", "422", "La hora no cabe en la jornada"],
        ],
        [2.25, 0.8, 3.55],
    )
    add_question(document, "Que demuestra el procedimiento", "Que la operacion critica se protege cerca de los datos. Aunque se conecten varios clientes o existan varios procesos de FastAPI, todos pasan por la misma transaccion de MySQL.")
    page_break(document)

    # Pagina 8
    add_heading(document, "7 Endpoints de la API")
    add_text(document, "Un endpoint es una direccion de la API para una operacion especifica. El metodo HTTP indica que tipo de operacion se desea realizar. GET consulta informacion, POST crea un recurso y PATCH modifica parcialmente un recurso existente.")
    add_table(
        document,
        ["Metodo y ruta", "Uso", "Respuesta importante"],
        [
            ["GET horarios id_doctor", "Consulta agenda para fecha y servicio", "Lista horas disponibles y ocupadas"],
            ["POST citas reservar", "Intenta confirmar una cita", "201 confirmada o 409 conflicto"],
            ["PATCH citas cancelar id_cita", "Cambia el estado de una cita", "Cita cancelada o aviso de que no es cancelable"],
            ["GET health", "Comprueba que la API esta encendida", "estado ok"],
        ],
        [1.75, 2.55, 2.3],
    )
    add_heading(document, "Consulta de horarios", 2)
    add_text(document, "La ruta de horarios recibe el identificador del medico, una fecha y el identificador del servicio. El servicio importa porque su duracion define el tamano de cada turno. Por ejemplo, con una consulta de treinta minutos y una jornada de 09:00 a 14:00 se generan bloques de media hora que no rebasan el final de la jornada.")
    add_heading(document, "Reserva y cancelacion", 2)
    add_text(document, "La reserva valida primero que existan paciente, medico y servicio. Luego revisa la jornada. Finalmente llama al procedimiento almacenado. La cancelacion ejecuta una actualizacion atomica: cambia el estado de Confirmada o Pendiente a Cancelada Paciente o Cancelada Medico y guarda la fecha de cancelacion.")
    add_text(document, "HTTP 409 es importante para la concurrencia. No es un fallo del programa: comunica que la solicitud estaba bien formada, pero no puede completarse porque el estado actual ya cambio. Flet responde mostrando que el horario acaba de ser reservado y vuelve a consultar disponibilidad.")
    add_heading(document, "Pregunta probable", 2)
    add_question(document, "Por que se usa PATCH para cancelar", "Porque no se borra fisicamente la cita. Solo se modifica su estado y se conserva el historial de cancelacion. PATCH representa una modificacion parcial del recurso existente.")
    page_break(document)

    # Pagina 9
    add_heading(document, "8 Configuracion demostracion y migracion")
    add_text(document, "El proyecto puede trasladarse a otra laptop, pero hay que separar el codigo de los datos locales. El codigo viaja como carpeta o archivo ZIP. La base debe crearse o importarse en MySQL de la laptop destino. El entorno virtual de Python no se recomienda copiar porque depende de la instalacion local de cada equipo.")
    add_table(
        document,
        ["Archivo o elemento", "Para que sirve", "Que hacer en otra laptop"],
        [
            ["db 01 instalar base corregida sql", "Crea tablas indices y procedimiento", "Ejecutar una vez en Workbench para una demo nueva"],
            ["db 02 cargar datos demo sql", "Agrega medico paciente servicio y horario", "Ejecutar despues del archivo 01"],
            ["archivo env", "Guarda host usuario y contraseña de MySQL", "Crear desde env example y cambiar la contraseña"],
            ["carpeta venv", "Dependencias instaladas de Python", "No copiar. Crear otra vez con PREPARAR PROYECTO"],
        ],
        [2.15, 2.4, 2.05],
    )
    add_heading(document, "Que se llena sola y que se captura", 2)
    add_text(document, "La estructura de las tablas no se llena sola. El script de datos demo agrega los registros necesarios para presentar. Durante el uso normal, Flet crea y cancela registros de la tabla citas. Si el sistema creciera, se agregarian pantallas administrativas o endpoints adicionales para registrar pacientes, doctores, servicios y horarios.")
    add_heading(document, "Secuencia de inicio", 2)
    add_bullet(document, "Instalar Python, MySQL Server y MySQL Workbench en la laptop de presentacion.")
    add_bullet(document, "Ejecutar el SQL 01 y luego el SQL 02 en Workbench.")
    add_bullet(document, "Ejecutar PREPARAR_PROYECTO.bat y escribir la contraseña correcta en .env.")
    add_bullet(document, "Abrir INICIAR_BACKEND.bat y no cerrar esa ventana.")
    add_bullet(document, "Abrir INICIAR_FLET.bat y realizar la prueba de consulta, reserva y cancelacion.")
    add_text(document, "Si ya existe una base llena que se desea conservar, no se debe ejecutar el archivo 01 porque reconstruye sistema_hospitalario. En ese caso se exporta estructura y datos desde MySQL Workbench y se importa el archivo resultante en la otra laptop.")
    page_break(document)

    # Pagina 10
    add_heading(document, "9 Preguntas para la defensa")
    add_question(document, "Que problema resuelve el proyecto", "Gestiona la disponibilidad de citas y evita que dos usuarios confirmen el mismo turno medico.")
    add_question(document, "Que es frontend", "La parte que usa la persona. Aqui es Flet con campos botones lista de horarios y mensajes.")
    add_question(document, "Que es backend", "La parte que recibe peticiones, aplica reglas y se comunica de forma controlada con MySQL. Aqui es FastAPI.")
    add_question(document, "Que es una API", "Un conjunto de rutas por las que una aplicacion puede solicitar operaciones a otra. En este proyecto Flet solicita operaciones a FastAPI mediante HTTP.")
    add_question(document, "Por que FastAPI y no conexion directa", "Porque centraliza validacion y evita exponer las credenciales y consultas de la base en la interfaz.")
    add_question(document, "Que es MySQL", "El gestor de base de datos que almacena los registros de manera permanente.")
    add_question(document, "Que es InnoDB", "El motor de MySQL que permite transacciones bloqueos por fila y llaves foraneas.")
    add_question(document, "Que es una transaccion", "Un conjunto de operaciones que se confirma completo con COMMIT o se deshace completo con ROLLBACK.")
    add_question(document, "Que es una condicion de carrera", "Un error posible cuando dos procesos leen el mismo recurso libre y lo modifican casi al mismo tiempo sin coordinacion.")
    add_question(document, "Como se evita la doble reserva", "sp_reservar_cita usa una transaccion y SELECT FOR UPDATE para bloquear la agenda del medico mientras decide e inserta la cita.")
    add_question(document, "Que representa el codigo 409", "Un conflicto. La solicitud estaba bien, pero el horario ya no se puede reservar porque otra operacion cambio la agenda.")
    add_question(document, "La consulta de horarios garantiza la reserva", "No. Solo orienta al usuario. La garantia final ocurre al reservar dentro de la transaccion de MySQL.")
    add_question(document, "Que ocurre al cancelar", "Se cambia el estado y se guarda la fecha de cancelacion. El historial se conserva y el procedimiento permite reutilizar ese horario.")
    add_heading(document, "Orden recomendado para la demostracion", 2)
    add_bullet(document, "Mostrar MySQL Workbench con el paciente doctor servicio y horario cargados.")
    add_bullet(document, "Abrir Flet, consultar horarios y reservar una hora disponible.")
    add_bullet(document, "Mostrar en Workbench que aparecio una fila nueva en citas.")
    add_bullet(document, "Intentar reservar la misma hora desde una segunda ejecucion o explicar que la API devuelve 409 si dos solicitudes compiten.")
    add_bullet(document, "Cancelar la cita y mostrar que el estado cambio sin borrar el registro.")

    return document


if __name__ == "__main__":
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
